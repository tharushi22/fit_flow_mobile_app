import docx
from docx.shared import Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOCX_PATH = r"D:\FitFlow\doc\IT3060_Lab06_FitFlow_Final_Report.docx"
GIT_URL = "https://github.com/tharushi22/fit_flow_mobile_app"

doc = docx.Document(DOCX_PATH)

# 1. Add GitHub Repository row to Cover Page Table (Table 0)
t0 = doc.tables[0]
new_row = t0.add_row()
new_row.cells[0].text = "GitHub Repository"
new_row.cells[1].text = GIT_URL

# Style the new row cells
for cell in new_row.cells:
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.name = "Segoe UI"
            run.font.size = Pt(10)

# Make GitHub URL bold blue
for p in new_row.cells[1].paragraphs:
    for run in p.runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(37, 99, 235)

print("Added GitHub Repository row to Cover Page (Table 0)")

# 2. Add GitHub Repository in Section 1 (Introduction)
for i, p in enumerate(doc.paragraphs):
    if "FitFlow is a fitness application redesigned" in p.text:
        # Add after this paragraph
        p_intro = doc.paragraphs[i+1] # paragraph after
        p_git = p_intro.insert_paragraph_before()
        p_git.paragraph_format.space_before = Pt(6)
        p_git.paragraph_format.space_after = Pt(6)
        
        r1 = p_git.add_run("Official GitHub Repository: ")
        r1.font.name = "Segoe UI"
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = RGBColor(30, 41, 59)
        
        r2 = p_git.add_run(GIT_URL)
        r2.font.name = "Segoe UI"
        r2.font.bold = True
        r2.font.size = Pt(10.5)
        r2.font.underline = True
        r2.font.color.rgb = RGBColor(37, 99, 235)
        print("Added GitHub link in Section 1 Introduction")
        break

# 3. Add GitHub Repository in Section 9 (Conclusion)
for i, p in enumerate(doc.paragraphs):
    if "FitFlow was successfully prepared as a signed Android release" in p.text:
        p_concl = doc.paragraphs[i]
        p_end = p_concl.insert_paragraph_before()
        p_end.paragraph_format.space_before = Pt(8)
        p_end.paragraph_format.space_after = Pt(4)
        
        r_repo = p_end.add_run("Project Codebase & Assets: ")
        r_repo.font.name = "Segoe UI"
        r_repo.font.bold = True
        r_repo.font.size = Pt(10)
        
        r_link = p_end.add_run(GIT_URL)
        r_link.font.name = "Segoe UI"
        r_link.font.underline = True
        r_link.font.color.rgb = RGBColor(37, 99, 235)
        print("Added GitHub link in Section 9 Conclusion")
        break

doc.save(DOCX_PATH)
print("Saved updated docx with GitHub link!")
